import base64
import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pypdf import PdfReader

from app.layout import build_document, format_birth_date, format_date, profile_url, years_of_experience
from app.models import CV

SAMPLE = json.loads((Path(__file__).parent / "sample_cv.json").read_text())


def _cv(template="international") -> dict:
    data = json.loads(json.dumps(SAMPLE))
    data["design"]["template"] = template
    return data


@pytest.mark.parametrize("template", ["international", "europass"])
def test_pdf_is_text_based_and_complete(client, template):
    r = client.post("/api/export/pdf", json=_cv(template))
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["x-page-count"] == "1"
    assert "Alex_Morgan_CV.pdf" in r.headers["content-disposition"]
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    for needle in ("Alex Morgan", "alex.morgan@example.com", "Northwind Fibre",
                   "cutting mean time to repair by 30%", "Python, Kafka", "German"):
        assert needle in text
    assert "Hidden section" not in text  # invisible sections are skipped
    if template == "europass":
        assert "WORK EXPERIENCE" in text and "Nationality" in text and "03/2021" in text
    else:
        assert "EXPERIENCE" in text and "Nationality" not in text and "Mar 2021" in text


@pytest.mark.parametrize("template", ["international", "europass"])
def test_docx_is_ats_safe(client, template):
    r = client.post("/api/export/docx", json=_cv(template))
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    assert len(doc.tables) == 0, "ATS-safe DOCX must not use tables"
    for s in doc.sections:  # nothing hidden in header/footer
        assert not any(p.text.strip() for p in s.header.paragraphs)
        assert not any(p.text.strip() for p in s.footer.paragraphs)
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        xml = z.read("word/document.xml").decode()
    assert "txbxContent" not in xml and "<w:drawing" not in xml  # no text boxes / images without a photo
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Alex Morgan" in text and "Northwind Fibre" in text
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings[0] in ("PROFESSIONAL SUMMARY", "PROFILE")
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) == 5
    assert "alex.morgan@example.com" in xml  # hyperlinked text still in the body


def test_custom_heading_and_order():
    data = _cv()
    data["sections"][1]["title"] = "Career History"
    data["sections"].insert(0, data["sections"].pop(3))  # skills first
    doc = build_document(CV.model_validate(data))
    assert [b.heading for b in doc.blocks][:2] == ["Skills", "Professional Summary"]
    assert "Career History" in [b.heading for b in doc.blocks]


def test_empty_cv_renders(client):
    r = client.post("/api/export/pdf", json={})
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    r = client.post("/api/export/docx", json={"sections": [{"type": "custom", "items": [{}]}]})
    assert r.status_code == 200


def test_unicode_text_survives(client):
    data = _cv()
    data["personal"]["fullName"] = "Zoë Łukasiewicz-Müller"
    r = client.post("/api/export/pdf", json=data)
    text = PdfReader(io.BytesIO(r.content)).pages[0].extract_text()
    assert "Zoë Łukasiewicz-Müller" in text


def test_validation_rejects_bad_template(client):
    data = _cv()
    data["design"]["template"] = "fancy-two-column"
    assert client.post("/api/export/pdf", json=data).status_code == 422


def test_dates_and_years():
    assert format_date("2021-03", "month_name") == "Mar 2021"
    assert format_date("2021-03", "numeric") == "03/2021"
    assert format_date("2021", "numeric") == "2021"
    assert format_date("Summer 2020", "numeric") == "Summer 2020"
    from datetime import date
    yrs = years_of_experience(CV.model_validate(SAMPLE), today=date(2026, 3, 1))
    assert 8.0 <= yrs <= 8.3


def test_meta_and_health(client):
    assert client.get("/api/health").json()["status"] == "ok"
    meta = client.get("/api/meta").json()
    assert {t["key"] for t in meta["templates"]} == {"international", "europass"}


def test_exports_require_login(anon):
    assert anon.post("/api/export/pdf", json=_cv()).status_code == 401
    assert anon.post("/api/export/docx", json=_cv()).status_code == 401


def _photo(w=300, h=500, mode="RGB") -> str:
    """A synthetic test image (no real person) as the data URL the editor sends."""
    im = Image.new(mode, (w, h), (40, 90, 160) if mode == "RGB" else (40, 90, 160, 128))
    buf = io.BytesIO()
    im.save(buf, "JPEG" if mode == "RGB" else "PNG")
    kind = "jpeg" if mode == "RGB" else "png"
    return f"data:image/{kind};base64," + base64.b64encode(buf.getvalue()).decode()


@pytest.mark.parametrize("template", ["international", "europass"])
def test_photo_in_both_formats(client, template):
    data = _cv(template)
    data["personal"]["photo"] = _photo()

    pdf = PdfReader(io.BytesIO(client.post("/api/export/pdf", json=data).content))
    assert len(pdf.pages) == 1 and len(pdf.pages[0].images) == 1
    text = pdf.pages[0].extract_text()
    assert text.index("Alex Morgan") < text.index("alex.morgan@example.com") < text.index("Northwind Fibre")

    r = client.post("/api/export/docx", json=data)
    doc = Document(io.BytesIO(r.content))
    assert len(doc.tables) == 0  # still single column, text not in a table
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        xml = z.read("word/document.xml").decode()
        assert any(n.startswith("word/media/") for n in z.namelist())
    assert xml.count("<w:drawing") == 1 and "<wp:anchor" in xml and "<wp:inline" not in xml  # floated, not inline
    assert 'descr="Photo of Alex Morgan"' in xml and 'w:clear="all"' in xml
    assert "txbxContent" not in xml
    for s in doc.sections:
        assert not s.header._element.xpath(".//w:drawing")
    assert doc.paragraphs[0].text == "Alex Morgan"


def test_photo_is_normalised_and_bad_photos_are_handled(client):
    doc = build_document(CV.model_validate({"personal": {"photo": _photo(800, 300, "RGBA")}}))
    with Image.open(io.BytesIO(doc.photo)) as im:  # landscape PNG with alpha -> 7:9 RGB JPEG
        assert im.format == "JPEG" and im.mode == "RGB" and im.size == (420, 540)

    data = _cv()
    data["personal"]["photo"] = "https://example.com/me.jpg"  # only embedded images are accepted
    assert client.post("/api/export/pdf", json=data).status_code == 422
    data["personal"]["photo"] = "data:image/jpeg;base64," + base64.b64encode(b"not an image").decode()
    r = client.post("/api/export/docx", json=data)  # undecodable: the CV still exports, without a photo
    assert r.status_code == 200 and b"word/media/" not in r.content


def test_birth_date_and_phone_code_only():
    assert format_birth_date("1990-03-15") == "15/03/1990"
    assert format_birth_date("15 March 1990") == "15 March 1990"  # older free text kept as typed
    data = _cv("europass")
    data["personal"]["dateOfBirth"] = "1990-03-15"
    data["personal"]["phone"] = "+964"  # a code picked, no number yet
    doc = build_document(CV.model_validate(data))
    assert ("Date of birth", "15/03/1990") in doc.details
    assert all(c.text != "+964" for c in doc.contacts)


def test_profile_links_point_at_the_profile():
    assert profile_url("linkedin", "alexmorgan") == "linkedin.com/in/alexmorgan"  # a bare handle
    assert profile_url("github", "@alex-morgan") == "github.com/alex-morgan"
    assert profile_url("linkedin", "https://www.linkedin.com/in/alexmorgan/") == "https://www.linkedin.com/in/alexmorgan/"
    assert profile_url("github", "alexmorgan.dev") == "alexmorgan.dev"  # a portfolio site
    # only the site's name: would link to its home page, so nothing is printed
    for site, value in [("linkedin", "LinkedIn"), ("linkedin", "linkedin.com"), ("linkedin", "https://www.linkedin.com/in/"),
                        ("github", "GitHub"), ("github", "www.github.com/"), ("github", "HTTPS://GitHub.com")]:
        assert profile_url(site, value) == "", value

    data = _cv()
    data["personal"].update(linkedin="linkedin.com", github="alexmorgan", website="")
    urls = [c.url for c in build_document(CV.model_validate(data)).contacts if c.url.startswith("https://")]
    assert urls == ["https://github.com/alexmorgan"]
