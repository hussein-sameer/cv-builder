"""ATS-safe DOCX renderer (python-docx).

Rules followed: single column, no tables, no text boxes, nothing in
headers/footers, real Word bullet lists, real "Heading 1" section headings,
standard fonts, contact details in the body.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Mm, Pt, RGBColor

from .layout import Block, Entry, Line, Link, RenderDoc

MUTED = RGBColor(0x44, 0x44, 0x44)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _style_font(style, name: str, size: float | None = None, bold: bool | None = None,
                color: RGBColor | None = None) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


# Elements that must come *after* w:pBdr inside w:pPr (OOXML schema order).
_PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _bottom_border(paragraph, hex_color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), hex_color.lstrip("#").upper())
    bdr.append(bottom)
    ppr.insert_element_before(bdr, *_PBDR_SUCCESSORS)


def _add_hyperlink(paragraph, text: str, url: str, color: RGBColor | None = None,
                   size: float | None = None) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if color is not None:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), str(color))
        rpr.append(c)
    if size is not None:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(round(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _run(paragraph, text: str, bold=False, italic=False, size=None, color=None):
    r = paragraph.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


class _Writer:
    def __init__(self, doc: RenderDoc):
        self.src = doc
        self.d = Document()
        self.fs = doc.font_size
        self.accent = _rgb(doc.accent)
        self._setup_page()
        self._setup_styles()

    # ----------------------------------------------------------------- setup
    def _setup_page(self) -> None:
        s = self.d.sections[0]
        if self.src.page_size == "Letter":
            s.page_width, s.page_height = Mm(215.9), Mm(279.4)
        else:
            s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = Mm(17)
        s.top_margin = s.bottom_margin = Mm(14)
        self.content_width = Emu(s.page_width - s.left_margin - s.right_margin)

    def _setup_styles(self) -> None:
        font, fs = self.src.font, self.fs
        normal = self.d.styles["Normal"]
        _style_font(normal, font, fs, color=RGBColor(0, 0, 0))
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = 1.05

        h1 = self.d.styles["Heading 1"]
        _style_font(h1, font, fs * 1.1, bold=True, color=self.accent)
        h1.font.italic = False
        pf = h1.paragraph_format
        pf.space_before = Pt(fs)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
        pf.line_spacing = 1.0

        lb = self.d.styles["List Bullet"]
        _style_font(lb, font, fs)
        lb.paragraph_format.space_after = Pt(0)
        lb.paragraph_format.line_spacing = 1.05

        cp = self.d.core_properties
        cp.title = f"{self.src.name} – CV" if self.src.name else "CV"
        cp.author = self.src.name or ""
        cp.subject = "Curriculum Vitae"
        cp.keywords = self.src.headline or ""

    # ----------------------------------------------------------------- parts
    def _p(self, style: str | None = None, keep_next=False, align=None, space_after: float | None = None):
        p = self.d.add_paragraph(style=style)
        if keep_next:
            p.paragraph_format.keep_with_next = True
        if align is not None:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def _link_or_text(self, p, link: Link, color=None, size: float | None = None) -> None:
        if link.url:
            _add_hyperlink(p, link.text, link.url, color, size)
        else:
            _run(p, link.text, color=color, size=size)

    def header(self) -> None:
        src = self.src
        align = WD_ALIGN_PARAGRAPH.CENTER if src.template.entry_style == "title_first" else WD_ALIGN_PARAGRAPH.LEFT
        if src.name:
            _run(self._p(align=align), src.name, bold=True, size=self.fs * 2.1)
        if src.headline:
            _run(self._p(align=align, space_after=2), src.headline, size=self.fs * 1.15, color=self.accent)
        if src.contacts:
            p = self._p(align=align)
            small = self.fs * 0.95
            for i, c in enumerate(src.contacts):
                if i:
                    _run(p, " | ", color=MUTED, size=small)
                self._link_or_text(p, c, MUTED, small)
        if src.details:
            p = self._p(align=align)
            for i, (k, v) in enumerate(src.details):
                if i:
                    _run(p, " | ", color=MUTED, size=self.fs * 0.95)
                _run(p, f"{k}: ", bold=True, color=MUTED, size=self.fs * 0.95)
                _run(p, v, color=MUTED, size=self.fs * 0.95)
        self._p(space_after=0)

    def heading(self, text: str) -> None:
        p = self.d.add_heading(text.upper(), level=1)
        _bottom_border(p, self.src.accent)

    def entry(self, e: Entry) -> None:
        dates_first = self.src.template.entry_style == "dates_first"
        paras = []
        if dates_first:
            meta = " | ".join(x for x in (e.dates, e.location) if x)
            if meta:
                p = self._p(keep_next=True)
                _run(p, meta, color=MUTED)
                paras.append(p)
            if e.title or e.subtitle:
                p = self._p(keep_next=True)
                if e.title:
                    _run(p, e.title, bold=True)
                if e.subtitle:
                    _run(p, (" – " if e.title else "") + e.subtitle)
                paras.append(p)
        else:
            p = self._p(keep_next=True)
            p.paragraph_format.tab_stops.add_tab_stop(self.content_width, WD_TAB_ALIGNMENT.RIGHT)
            _run(p, e.title or e.subtitle, bold=bool(e.title))
            if e.dates:
                _run(p, "\t" + e.dates)
            paras.append(p)
            org_loc = " – ".join(x for x in (e.subtitle, e.location) if x)
            if e.title and org_loc:
                p = self._p(keep_next=True)
                _run(p, org_loc, color=MUTED)
                paras.append(p)

        if e.link:
            p = self._p(keep_next=True)
            self._link_or_text(p, e.link, self.accent)
            paras.append(p)
        for x in e.extra_lines:
            p = self._p()
            _run(p, x)
            paras.append(p)
        for x in e.description:
            p = self._p()
            _run(p, x)
            paras.append(p)
        for b in e.bullets:
            p = self._p(style="List Bullet")
            _run(p, b)
            paras.append(p)
        if e.tags:
            p = self._p()
            _run(p, f"{e.tags_label}: ", bold=True)
            _run(p, e.tags)
            paras.append(p)
        if paras:
            paras[-1].paragraph_format.space_after = Pt(self.fs * 0.6)

    def line(self, ln: Line) -> None:
        p = self._p(space_after=2 if not ln.sub else 0)
        _run(p, ln.bold, bold=True)
        if ln.rest:
            _run(p, f" – {ln.rest}")
        if ln.link:
            _run(p, " | ")
            self._link_or_text(p, ln.link, self.accent)
        if ln.sub:
            sp = self._p(space_after=2)
            sp.paragraph_format.left_indent = Pt(8)
            _run(sp, ln.sub, size=self.fs * 0.95, color=MUTED)

    def block(self, b: Block) -> None:
        self.heading(b.heading)
        if b.kind == "paragraphs":
            for i, para in enumerate(b.paragraphs):
                _run(self._p(space_after=3 if i < len(b.paragraphs) - 1 else 0), para)
        elif b.kind == "entries":
            for e in b.entries:
                self.entry(e)
        elif b.kind == "pairs":
            for label, value in b.pairs:
                p = self._p(space_after=1.5)
                if label and value:
                    _run(p, f"{label}: ", bold=True)
                    _run(p, value)
                else:
                    _run(p, label or value)
        elif b.kind == "lines":
            for ln in b.lines:
                self.line(ln)

    def build(self) -> bytes:
        # python-docx starts with one empty paragraph in some templates; remove it
        body = self.d.element.body
        for p in list(body.iterchildren(qn("w:p"))):
            if not p.xpath(".//w:t"):
                body.remove(p)
        self.header()
        for b in self.src.blocks:
            self.block(b)
        buf = io.BytesIO()
        self.d.save(buf)
        return buf.getvalue()


def render_docx(doc: RenderDoc) -> bytes:
    return _Writer(doc).build()
